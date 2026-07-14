import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# =====================================================================
# ESTILOS
# =====================================================================
header_font = Font(bold=True, color="FFFFFF", size=10)
header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
wrap = Alignment(wrap_text=True, vertical="top")
wrap_center = Alignment(wrap_text=True, vertical="top", horizontal="center")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

tier_fills = {
    "TIER 1": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "TIER 2": PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid"),
    "TIER 3": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
}
tier_fonts = {
    "TIER 1": Font(bold=True, size=10, color="006100"),
    "TIER 2": Font(bold=False, size=10, color="1F3864"),
    "TIER 3": Font(bold=False, size=10, color="595959"),
}

# =====================================================================
# HOJA 1: DIRECTORIO MAESTRO INTEGRADO
# =====================================================================
ws = wb.active
ws.title = "Directorio Maestro BATEV 2026"

headers = [
    "Tier",                          # 1
    "Empresa",                       # 2
    "Persona / Rol",                 # 3
    "Teléfono",                      # 4
    "Email",                         # 5
    "Web / IG",                      # 6
    "Stand",                         # 7
    "Rubro",                         # 8
    "Qué vende",                     # 9
    "Nivel Sponsor",                 # 10
    "Prioridad Comercial",           # 11
    "Estado del dato",               # 12
    "Dolor / Insight principal",     # 13
    "Sistema actual",                # 14
    "Quién decide",                  # 15
    "Objeción / Contexto",           # 16
    "Gancho / Ángulo AureaHub",      # 17
    "Qué le serviría (auditoría web)", # 18
    "Próximo paso",                  # 19
    "Contactado por",                # 20
    "Obs. contacto WhatsApp",        # 21
    "Fuente",                        # 22
]

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = thin_border

# =====================================================================
# DATOS INTEGRADOS — Ordenados por TIER
# =====================================================================
# Tier 1: Contactos e Insights (ida y vuelta real)
# Tier 2: Plantilla de contactos (teléfono/email, interacción en feria)
# Tier 3: Solo scrape web (sin contacto directo)

rows_data = [
    # ===== TIER 1: CONTACTOS E INSIGHTS (conversación real) =====
    {
        "tier": "TIER 1", "empresa": "Parrillas de acero (posible Pavir)",
        "persona": "(confirmar — tarjeta sin nombre claro)", "telefono": "+54 9 232 331 3208 / 02323-430270",
        "email": "", "web": "pavir.com.ar", "stand": "E-24", "rubro": "Parrillas / Aberturas",
        "que_vende": "3 marcas, 3 públicos: Camargo (cliente final, venta lenta), Bosca (100% mayorista), Napoleón (mayorista). Canales separados.",
        "sponsor": "Sponsor Diamond", "prioridad": "Alta",
        "estado": "Insight rico",
        "dolor": "WhatsApp borraba contexto cada 24h: perdían historial, no sabían en qué estadío estaba la venta, leads que se esfumaban. Sin calificación por estadío.",
        "sistema": "WhatsApp sin memoria. Agencia de marketing interiorizada.",
        "decide": "Presidente (si es Pavir)",
        "objecion": "Permisos de redes complejos por licenciatarios (Bosca, Napoleón).",
        "gancho": "Memoria y contexto persistente + calificación por estadío. Embudos y Smart Tags por marca: Camargo = venta, Bosca = postventa.",
        "auditoria": "Sin formulario de cotización ni chatbot; podría automatizar consultas de WhatsApp y capturar leads con un embudo de ventas.",
        "proximo": "Reunión 1ra sem julio. Confirmar de quién es la tarjeta.",
        "contactado": "", "obs_wa": "", "fuente": "Audio + notas feria",
    },
    {
        "tier": "TIER 1", "empresa": "PAVIR Puertas",
        "persona": "Daniel Paissan (a confirmar). Presidente", "telefono": "+54 9 232 331 3208 / 02323-430270",
        "email": "dpaissan@pavir.com.ar", "web": "pavir.com.ar", "stand": "E-24", "rubro": "Aberturas",
        "que_vende": "Puertas pivotantes de alta prestación en acero, PVC y aluminio.",
        "sponsor": "Sponsor Diamond", "prioridad": "Alta",
        "estado": "Insight rico (vinculado a Parrillas)",
        "dolor": "(ver fila Parrillas — misma empresa probable)",
        "sistema": "", "decide": "Presidente", "objecion": "", "gancho": "",
        "auditoria": "Sin formulario de cotización ni chatbot; podría automatizar consultas de WhatsApp y capturar leads con un embudo de ventas.",
        "proximo": "Reunión 1ra sem julio.", "contactado": "", "obs_wa": "", "fuente": "Tarjeta feria + scrape",
    },
    {
        "tier": "TIER 1", "empresa": "Emotorart (nombre a confirmar)",
        "persona": "(sin datos de contacto aún)", "telefono": "", "email": "", "web": "", "stand": "", "rubro": "Por confirmar",
        "que_vende": "Por confirmar.",
        "sponsor": "", "prioridad": "Alta y URGENTE",
        "estado": "Sin contacto",
        "dolor": "Está por cerrar un desarrollo a medida con OTRA agencia (competencia directa).",
        "sistema": "Evaluando desarrollo a medida.",
        "decide": "", "objecion": "Competencia: se va a un desarrollo a medida.",
        "gancho": "Producto ya andando vs desarrollo a medida que se cae. La analogía auto hecho vs a medida ya funcionó.",
        "auditoria": "URGENTE: conseguir datos y contactar antes de que cierre con la competencia.",
        "proximo": "Contactar YA, antes de que cierre. Conseguir datos de contacto.",
        "contactado": "", "obs_wa": "", "fuente": "Audio feria",
    },
    {
        "tier": "TIER 1", "empresa": "Macoser",
        "persona": "Horacio Lagos. Dueño / senior (70 años)", "telefono": "",
        "email": "Horacio.lagos@macoser.com.ar", "web": "", "stand": "", "rubro": "Hornos / Industria",
        "que_vende": "Hornos (insumos / industria).",
        "sponsor": "", "prioridad": "Alta",
        "estado": "Insight parcial",
        "dolor": "Tiene ERP y le falta la capa de ventas.",
        "sistema": "ERP, sin CRM.",
        "decide": "Horacio (parece decisor).",
        "objecion": "Sin objeción, interesado y con ganas.",
        "gancho": "Capa comercial sobre el ERP sin tocarlo. Cadencia suave, no atosigar (tranquilo, 70 años).",
        "auditoria": "Necesita CRM liviano que se monte sobre su ERP existente; automatizar seguimiento comercial sin reemplazar lo que ya tiene.",
        "proximo": "Reunión 1ra sem julio. Trato tranquilo.",
        "contactado": "", "obs_wa": "", "fuente": "Audio feria",
    },
    {
        "tier": "TIER 1", "empresa": "American Tint",
        "persona": "Persona de marketing (nombre a confirmar)", "telefono": "+54 9 11 6155 0006",
        "email": "info@americantint.ar", "web": "americantint.ar · IG @americantint.ar", "stand": "", "rubro": "Films / Láminas",
        "que_vende": "Films para vidrios y gráfica. Empresa nueva, pocas ventas todavía.",
        "sponsor": "Sponsor Platinum", "prioridad": "Media",
        "estado": "Insight parcial",
        "dolor": "Quieren armar la base de datos para nutrir info y futuras estrategias comerciales. Demanda BAJA hoy.",
        "sistema": "Nada aún.",
        "decide": "Marketing (confirmar si decide).",
        "objecion": "Empresa nueva, sin volumen. Confirmar si necesitan generar demanda (Bardo) u ordenar (Prometheo).",
        "gancho": "Armar la base ordenada desde el día cero para escalar.",
        "auditoria": "Sin WhatsApp ni formulario de cotización visibles; un chatbot podría convertir el catálogo en consultas calificadas.",
        "proximo": "Hablar en 2 semanas.",
        "contactado": "", "obs_wa": "", "fuente": "Audio feria + scrape",
    },
    {
        "tier": "TIER 1", "empresa": "Lindall S.A. (capsulasaura)",
        "persona": "Rodolfo Ortiz. Gerencia de Ventas", "telefono": "+54 351 211 4832",
        "email": "rodolfoneris@gmail.com (verificar)", "web": "capsulasaura.com · Córdoba Capital", "stand": "", "rubro": "Por confirmar (naming)",
        "que_vende": "Por confirmar (naming Lindall vs capsulasaura).",
        "sponsor": "", "prioridad": "Por definir",
        "estado": "Solo tarjeta",
        "dolor": "", "sistema": "", "decide": "Gerencia de Ventas (buen rol, decide).", "objecion": "", "gancho": "",
        "auditoria": "Ventas 100% manuales vía WhatsApp sin chatbot ni CRM; un bot de WhatsApp podría calificar leads y agendar automáticamente.",
        "proximo": "Reunión 1ra sem julio. Aclarar el naming.",
        "contactado": "", "obs_wa": "", "fuente": "Tarjeta feria",
    },
    {
        "tier": "TIER 1", "empresa": "GU Herrajes S.A.",
        "persona": "Gastón Diaz. Técnica", "telefono": "+54 11 3096-5886",
        "email": "tecnica@gu-herrajes.com.ar", "web": "gu-herrajes.com.ar · Saenz Peña", "stand": "", "rubro": "Herrajes aberturas",
        "que_vende": "Herrajes para aberturas (GU, BKS, FERCO).",
        "sponsor": "", "prioridad": "Por definir",
        "estado": "Solo tarjeta",
        "dolor": "", "sistema": "", "decide": "Técnica (ojo Authority, no comercial).", "objecion": "", "gancho": "",
        "auditoria": "Tiene formulario de cotización pero sin WhatsApp, chatbot ni CRM; falta automatizar el seguimiento de cada solicitud.",
        "proximo": "Reunión 1ra sem julio.",
        "contactado": "", "obs_wa": "", "fuente": "Tarjeta feria",
    },
    {
        "tier": "TIER 1", "empresa": "Cristales Sáenz",
        "persona": "(sin nombre en la tarjeta)", "telefono": "11 3452 3008",
        "email": "cristalessaenz@outlook.com", "web": "cristales-saenz.com.ar · Lomas del Mirador", "stand": "", "rubro": "Cristales / DVH",
        "que_vende": "Cristales laminados / DVH.",
        "sponsor": "", "prioridad": "Por definir",
        "estado": "Solo tarjeta",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Solo WhatsApp y datos de contacto, sin formulario, catálogo online ni chatbot para calificar consultas de DVH.",
        "proximo": "Reunión 1ra sem julio.",
        "contactado": "", "obs_wa": "", "fuente": "Tarjeta feria",
    },
    {
        "tier": "TIER 1", "empresa": "Asothella",
        "persona": "Diego Palacios", "telefono": "011-5642-5365 / 011-4652-2305",
        "email": "infotejas@tejas-asothella.com", "web": "asothella.com.ar · La Tablada", "stand": "", "rubro": "Techos / Tejas",
        "que_vende": "Tejas metálicas.",
        "sponsor": "", "prioridad": "Por definir",
        "estado": "Solo tarjeta",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Tiene WhatsApp y carrito, pero sin chat en vivo ni formulario propio; pierde leads que no llaman por WhatsApp.",
        "proximo": "Reunión 1ra sem julio.",
        "contactado": "", "obs_wa": "Fijo (verificar WhatsApp). Sin respuesta en la feria.", "fuente": "Tarjeta feria + plantilla contactos",
    },
    {
        "tier": "TIER 1", "empresa": "Rheem",
        "persona": "Equipo de Marketing. Ya era contacto previo (Pablo Malcovich)", "telefono": "0810-888-6060",
        "email": "comunicaciones@rheem.com.ar", "web": "rheem.com.ar · IG @rheem.argentina", "stand": "D-16", "rubro": "Agua Caliente",
        "que_vende": "Calefones, calderas SmartHeatDuo, termotanques residenciales/comerciales, híbridos.",
        "sponsor": "", "prioridad": "Por definir",
        "estado": "Solo tarjeta",
        "dolor": "", "sistema": "", "decide": "Marketing (ojo Authority).", "objecion": "", "gancho": "",
        "auditoria": "Solo ofrece teléfono y tienda online, sin WhatsApp ni chatbot; falta captura de leads y seguimiento tipo CRM.",
        "proximo": "Reunión 1ra sem julio. Ya era contacto previo.",
        "contactado": "", "obs_wa": "", "fuente": "Tarjeta feria + scrape",
    },

    # ===== TIER 2: PLANTILLA CONTACTOS (interacción en feria, teléfono/email) =====
    {
        "tier": "TIER 2", "empresa": "ALUWIND",
        "persona": "", "telefono": "5493516248155",
        "email": "aluwind12@gmail.com", "web": "aluwind.com.ar", "stand": "", "rubro": "Aluminio",
        "que_vende": "Perfiles y barandas de aluminio.",
        "sponsor": "Sponsor Platinum", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Fuerte en WhatsApp y calculadoras, pero sin chatbot ni CRM evidente; podría automatizar seguimiento de pedidos y cotizaciones.",
        "proximo": "", "contactado": "",
        "obs_wa": "Movil/WhatsApp. IA que contesta pero no admite ser IA.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "MUCHTEK",
        "persona": "", "telefono": "5491172214600",
        "email": "perfiles@muchtek.com", "web": "muchtek.com", "stand": "220 m²", "rubro": "Aberturas PVC",
        "que_vende": "Productos arquitectónicos de PVC: línea Evolution, Advance, Efficient, In&Out Design.",
        "sponsor": "Main Sponsor", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Depende de contacto directo sin chatbot ni formulario de cotización; un asistente IA agilizaría consultas técnicas de perfiles.",
        "proximo": "", "contactado": "",
        "obs_wa": "Movil/WhatsApp. Respondió persona ~3 hs después.", "fuente": "Plantilla contactos feria + scrape",
    },
    {
        "tier": "TIER 2", "empresa": "INOMAX",
        "persona": "", "telefono": "5493515997099",
        "email": "info@herrajesinomax.com.ar", "web": "herrajesinomax.com.ar", "stand": "", "rubro": "Herrajes",
        "que_vende": "Herrajes para vidrio. Venta exclusiva a profesionales.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Tienda B2B funcional pero sin chat en vivo ni formulario de contacto; un chatbot de WhatsApp agilizaría cotizaciones.",
        "proximo": "", "contactado": "",
        "obs_wa": "Movil/WhatsApp. Automático + humano en horario.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "WINTECH ARG",
        "persona": "", "telefono": "5493512326203",
        "email": "contacto@wintech.com.ar", "web": "wintech.com.ar", "stand": "", "rubro": "Aberturas PVC",
        "que_vende": "Perfiles premium PVC para aberturas.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Solo WhatsApp y teléfono, sin formulario ni captura de leads; un embudo con CRM automatizaría el contacto con la red de fabricantes.",
        "proximo": "", "contactado": "",
        "obs_wa": "Movil/WhatsApp. Humano horas después, insistió al otro día.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "GRUPO ELEKTRIM",
        "persona": "", "telefono": "541152735050",
        "email": "contacto@elektrim.com.ar", "web": "elektrim.com.ar", "stand": "", "rubro": "Bombas de agua",
        "que_vende": "Bombas de agua para el hogar.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "No tiene botón de WhatsApp ni chatbot; formulario de Google suelto sin integración a CRM para seguimiento de leads.",
        "proximo": "", "contactado": "",
        "obs_wa": "Fijo (verificar WhatsApp). Contestó humano al otro día.", "fuente": "Plantilla contactos feria + scrape",
    },
    {
        "tier": "TIER 2", "empresa": "ECOSAN",
        "persona": "", "telefono": "541148460995",
        "email": "info@ecosan.com.ar", "web": "ecosan.com.ar", "stand": "", "rubro": "Construcción / Proyectos",
        "que_vende": "Proyectos residenciales, comerciales e industriales.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Buen formulario y WhatsApp, pero sin chat en vivo ni agendamiento online; un chatbot IA precalificaría consultas de proyectos.",
        "proximo": "", "contactado": "",
        "obs_wa": "Fijo. Automático compartiendo WhatsApp directo.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "FONAC",
        "persona": "", "telefono": "5491153193065",
        "email": "acustica@decibel.com.ar", "web": "", "stand": "", "rubro": "Aislación acústica",
        "que_vende": "Soluciones acústicas para espacios.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Nicho acústico especializado; un catálogo interactivo con calculadora de aislación + captura de leads sería diferencial.",
        "proximo": "", "contactado": "",
        "obs_wa": "Movil/WhatsApp. Automático compartiendo WhatsApp directo.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "VERATECH / Fenstech",
        "persona": "", "telefono": "5491135817515",
        "email": "", "web": "fenstech.com (CAÍDO)", "stand": "", "rubro": "Aberturas PVC",
        "que_vende": "Venta general de productos del rubro construcción.",
        "sponsor": "Sponsor Diamond", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "ALERTA: dominio fenstech.com caído y en venta; están perdiendo toda oportunidad digital. Necesitan web + CRM urgente.",
        "proximo": "", "contactado": "",
        "obs_wa": "Movil/WhatsApp. VERIFICAR: figura como 'Fenstech Veratec'. Automático largo.", "fuente": "Plantilla contactos feria + scrape",
    },
    {
        "tier": "TIER 2", "empresa": "ANSAL",
        "persona": "", "telefono": "548102226725",
        "email": "ansal@ansal.com.ar", "web": "ansal.com.ar", "stand": "C-35", "rubro": "Climatización",
        "que_vende": "Climatización GREE: AC + piso radiante en un equipo; calderas, VRF, splits.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Catálogo B2B extenso pero sin WhatsApp, chatbot ni formulario online; ideal para un embudo de cotizaciones automatizado.",
        "proximo": "", "contactado": "",
        "obs_wa": "Línea 0810 (no WhatsApp). Automático compartiendo WA directo.", "fuente": "Plantilla contactos feria + scrape",
    },
    {
        "tier": "TIER 2", "empresa": "GENERAC",
        "persona": "", "telefono": "541136403895",
        "email": "generac@generac.com.ar", "web": "", "stand": "", "rubro": "Generadores",
        "que_vende": "Generadores.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Marca de generadores sin presencia web local evidente; un sitio con cotizador + WhatsApp captaría leads del sector construcción.",
        "proximo": "", "contactado": "",
        "obs_wa": "Fijo. Automático compartiendo WhatsApp directo.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "ARNEG",
        "persona": "", "telefono": "543414106100",
        "email": "ventas@arneg.com.ar", "web": "", "stand": "", "rubro": "Refrigeración comercial",
        "que_vende": "Equipamiento de refrigeración / construcción.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Refrigeración comercial: un configurador de equipos online + CRM para postventa diferencia mucho en este rubro.",
        "proximo": "", "contactado": "",
        "obs_wa": "Fijo. Automático compartiendo mails directo.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "ANUM",
        "persona": "", "telefono": "541172407777",
        "email": "", "web": "", "stand": "", "rubro": "Materiales construcción",
        "que_vende": "Importador de materiales de construcción.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Importador sin web evidente; necesita catálogo digital B2B con stock y precios para distribuidores + seguimiento CRM.",
        "proximo": "", "contactado": "",
        "obs_wa": "Fijo. Sin respuesta en la feria.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "BRIMAX",
        "persona": "", "telefono": "543416829886",
        "email": "info@brimaxargentina.com.ar", "web": "", "stand": "H-1", "rubro": "Ladrillos / HCCA",
        "que_vende": "Ladrillos y paneles de HCCA.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Fabricante de HCCA: un calculador de materiales por m² + WhatsApp bot para distribuidores captaría más consultas calificadas.",
        "proximo": "", "contactado": "",
        "obs_wa": "Fijo. Sin respuesta en la feria.", "fuente": "Plantilla contactos feria + scrape",
    },
    {
        "tier": "TIER 2", "empresa": "PRAMAX / PRAMAC",
        "persona": "", "telefono": "5491160938313",
        "email": "ventas@pramac.com.ar", "web": "", "stand": "", "rubro": "Generadores",
        "que_vende": "Generadores.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Marca de generadores premium; un cotizador por potencia requerida + CRM de seguimiento postventa agregaría valor inmediato.",
        "proximo": "", "contactado": "",
        "obs_wa": "Movil/WhatsApp. Sin respuesta en la feria.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "HIDROZONO",
        "persona": "", "telefono": "541155780606",
        "email": "administracion@hidrozono.com.ar", "web": "", "stand": "", "rubro": "Spa / Sauna",
        "que_vende": "Productos de spa y sauna.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Nicho spa/sauna premium: un catálogo visual con configurador de espacios + captura de leads por proyecto sería diferencial.",
        "proximo": "", "contactado": "",
        "obs_wa": "Fijo. Sin respuesta en la feria.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "MOZART",
        "persona": "", "telefono": "5491157814444",
        "email": "", "web": "", "stand": "", "rubro": "Grifería",
        "que_vende": "Grifería.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Grifería sin presencia web; necesita catálogo con filtros por línea + embudo de cotización para arquitectos y plomeros.",
        "proximo": "", "contactado": "",
        "obs_wa": "Movil/WhatsApp. Sin respuesta en la feria.", "fuente": "Plantilla contactos feria",
    },
    {
        "tier": "TIER 2", "empresa": "SAULEDA TEXTILES",
        "persona": "", "telefono": "(empresa española, sin línea local AR)",
        "email": "sauleda@sauleda.com", "web": "", "stand": "", "rubro": "Textiles",
        "que_vende": "Textiles.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío (sin teléfono)",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Empresa española sin presencia local; necesitaría landing AR con distribuidores + WhatsApp local para captar mercado argentino.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Plantilla contactos feria (sin tel)",
    },
    {
        "tier": "TIER 2", "empresa": "CONSUL STEEL",
        "persona": "", "telefono": "",
        "email": "consulsteel@adbarbieri.com", "web": "adbarbieri.com", "stand": "", "rubro": "Steel framing",
        "que_vende": "Academia de steel frame.",
        "sponsor": "Sponsor Platinum (Barbieri)", "prioridad": "", "estado": "Lead frío (sin teléfono)",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Academia de steel frame sin e-commerce ni WhatsApp; un funnel de inscripción + chatbot de consultas técnicas atraería más alumnos.",
        "proximo": "", "contactado": "", "obs_wa": "No publica teléfono; contacto por mail.", "fuente": "Plantilla contactos feria (sin tel)",
    },
    {
        "tier": "TIER 2", "empresa": "PEDROLLO",
        "persona": "", "telefono": "(sin nro oficial único en AR)",
        "email": "", "web": "", "stand": "", "rubro": "Bombas de agua",
        "que_vende": "Motores/bombas de agua.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío (sin teléfono)",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Marca italiana vía distribuidores AR; un directorio de distribuidores con geo-localizador + CRM centralizado optimizaría su canal.",
        "proximo": "", "contactado": "", "obs_wa": "Sin nro. oficial — vía distribuidores.", "fuente": "Plantilla contactos feria (sin tel)",
    },
    {
        "tier": "TIER 2", "empresa": "DUPLEX",
        "persona": "", "telefono": "",
        "email": "", "web": "IG @duplexargentina", "stand": "", "rubro": "Limpieza",
        "que_vende": "Productos de limpieza de hogar.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío (sin teléfono)",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Solo presencia en Instagram; necesita web + carrito o WhatsApp commerce para convertir seguidores en compradores.",
        "proximo": "", "contactado": "", "obs_wa": "Sin nro. público — IG @duplexargentina.", "fuente": "Plantilla contactos feria (sin tel)",
    },
    {
        "tier": "TIER 2", "empresa": "TOP FIBER",
        "persona": "", "telefono": "",
        "email": "", "web": "IG @topfiberbr", "stand": "", "rubro": "(sin dato)",
        "que_vende": "(sin dato).",
        "sponsor": "", "prioridad": "", "estado": "Lead frío (sin teléfono)",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Marca brasileña sin presencia argentina; si buscan entrar al mercado AR necesitan landing local + distribuidores.",
        "proximo": "", "contactado": "", "obs_wa": "Marca brasileña — IG @topfiberbr.", "fuente": "Plantilla contactos feria (sin tel)",
    },
    {
        "tier": "TIER 2", "empresa": "ELEKTA",
        "persona": "", "telefono": "",
        "email": "", "web": "IG @elektapuertas", "stand": "", "rubro": "Aberturas",
        "que_vende": "Puertas.",
        "sponsor": "", "prioridad": "", "estado": "Lead frío (sin teléfono)",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Solo Instagram, sin web ni WhatsApp público; necesita embudo básico para convertir consultas de IG en cotizaciones.",
        "proximo": "", "contactado": "", "obs_wa": "Sin nro. público — IG @elektapuertas.", "fuente": "Plantilla contactos feria (sin tel)",
    },

    # ===== TIER 3: SOLO SCRAPE WEB (sin contacto directo) =====
    {
        "tier": "TIER 3", "empresa": "Clarín ARQ", "persona": "", "telefono": "", "email": "", "web": "", "stand": "",
        "rubro": "Medios / Arquitectura", "que_vende": "Suplemento de arquitectura del diario Clarín.",
        "sponsor": "Main Sponsor", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Medio de prensa, no prospecto directo; útil para pauta publicitaria o PR de AureaHub.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Motorarg", "persona": "", "telefono": "", "email": "", "web": "motorarg.com", "stand": "",
        "rubro": "Herramientas / Maquinaria", "que_vende": "Herramientas y equipos para construcción.",
        "sponsor": "Sponsor Diamond", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Marca de herramientas con red de distribuidores; un portal B2B con stock y pedidos + CRM de distribuidores sería su upgrade.",
        "proximo": "", "contactado": "Valen", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Sinteplast", "persona": "", "telefono": "", "email": "", "web": "sinteplast.com.ar", "stand": "",
        "rubro": "Pinturas y Revestimientos", "que_vende": "Pinturas, impermeabilizantes, revestimientos.",
        "sponsor": "Sponsor Diamond", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Marca fuerte con herramientas virtuales, pero sin WhatsApp ni chat; deriva todo a puntos de venta sin automatizar el paso intermedio.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Vann", "persona": "", "telefono": "", "email": "", "web": "vann.com.ar", "stand": "",
        "rubro": "Griferías / Sanitarios", "que_vende": "Griferías y accesorios para baño y cocina.",
        "sponsor": "Sponsor Diamond", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Grifería premium; un showroom virtual 3D + embudo de cotización para arquitectos sería diferenciador en el rubro.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "A.D. Barbieri", "persona": "", "telefono": "", "email": "", "web": "barbieri.com.ar → adbarbieri.com", "stand": "",
        "rubro": "Construcción en seco", "que_vende": "Sistemas constructivos, Consulsteel, Barbieri|Deceuninck.",
        "sponsor": "Sponsor Platinum", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "E-commerce funcional pero sin WhatsApp ni chat en la home; un chatbot de WhatsApp eliminaría fricción en el contacto inmediato.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Patagonia Flooring", "persona": "", "telefono": "", "email": "", "web": "patagoniafloorings.com (CAÍDO)", "stand": "",
        "rubro": "Pisos", "que_vende": "Pisos de madera, vinílicos y laminados.",
        "sponsor": "Sponsor Platinum", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "ALERTA: dominio no resuelve (caído). Necesitan reactivar web urgente; están invisible digitalmente.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Wagg", "persona": "", "telefono": "", "email": "", "web": "wagg.com.ar", "stand": "E-21",
        "rubro": "Arquitectura Textil", "que_vende": "Cielorrasos Barrisol, cubiertas y fachadas textiles.",
        "sponsor": "Sponsor Platinum", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Tiene WhatsApp y newsletter, pero sin formulario de cotización ni CRM; falta funnel que convierta portfolio en pedidos.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Aluplexa", "persona": "", "telefono": "", "email": "", "web": "aluplexa.com", "stand": "",
        "rubro": "Aberturas / Persianas", "que_vende": "Persianas de grandes dimensiones, mosquiteras.",
        "sponsor": "Sponsor Gold", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Nicho persianas oversize; un cotizador por medidas + bot de WhatsApp calificaría consultas de arquitectos automáticamente.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Ensa Ascensores", "persona": "", "telefono": "", "email": "", "web": "ensaascensores.com", "stand": "",
        "rubro": "Ascensores", "que_vende": "Ascensores residenciales PVE52/37/30 (instalación 1 día).",
        "sponsor": "Sponsor Gold", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Producto diferencial (instala en 1 día); un landing con calculadora de viabilidad + captura de lead caliente sería muy efectivo.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Grupo LTN", "persona": "", "telefono": "", "email": "", "web": "", "stand": "",
        "rubro": "Soluciones constructivas", "que_vende": "Soluciones constructivas integrales.",
        "sponsor": "Sponsor Gold", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Sin web pública evidente; necesitan presencia digital completa (web + CRM + WhatsApp) para capitalizar su nivel de sponsor.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Puertas Brandsen", "persona": "", "telefono": "", "email": "", "web": "", "stand": "",
        "rubro": "Puertas", "que_vende": "Puertas para construcción.",
        "sponsor": "Sponsor Gold", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Fabricante de puertas sin web visible; un catálogo digital con medidas estándar y a medida + WhatsApp bot captaría distribuidores.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Home Access Argentina", "persona": "", "telefono": "", "email": "", "web": "", "stand": "",
        "rubro": "Accesibilidad", "que_vende": "Soluciones de acceso para el hogar.",
        "sponsor": "Sponsor Gold", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Nicho accesibilidad con poca competencia digital; un embudo educativo (normativa + productos) + CRM captaría leads de arquitectos.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Liquitech", "persona": "", "telefono": "", "email": "", "web": "liquitech.com.ar", "stand": "A-26",
        "rubro": "Pinturas / Impermeabilizantes", "que_vende": "AquaEpoxi, Cauchogoma (impermeabilizantes).",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Producto de renovación sin obra; un funnel educativo (antes/después) + WhatsApp bot convertiría curiosos en compradores.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Coef Track", "persona": "", "telefono": "", "email": "", "web": "coeftrack.com", "stand": "A-17",
        "rubro": "Software Construcción", "que_vende": "Sistema de gestión de personal de obra.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "SaaS de construcción: compañero natural de AureaHub, posible partnership más que cliente.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Finnegans GO", "persona": "", "telefono": "", "email": "", "web": "finneg.com/ar/", "stand": "C-28",
        "rubro": "Software Construcción", "que_vende": "Gestión de obras en la nube.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "ERP de construcción establecido; posible integración CRM + WhatsApp como complemento a su plataforma.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "IDERO", "persona": "", "telefono": "", "email": "", "web": "ideroarquitectura.com.ar", "stand": "F-7a",
        "rubro": "Construcción Modular", "que_vende": "Lodges premium modulares (18-67 m²).",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Sin WhatsApp, chat ni formulario; un funnel con captura de leads calificaría prospectos de lodges premium.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Almapiedra", "persona": "", "telefono": "", "email": "", "web": "almapiedra.com.ar", "stand": "H-11",
        "rubro": "Restauración / Nanotecnología", "que_vende": "Restauración patrimonial, nanotecnología Tecnan.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Sin WhatsApp ni chat; vende productos técnicos sin tienda online. Un chatbot de ventas destrabaría consultas técnicas.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Cambre", "persona": "", "telefono": "", "email": "", "web": "cambre.com.ar", "stand": "E-11a",
        "rubro": "Electricidad / Domótica", "que_vende": "Domótica Cambre Live Home, electricidad.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Ya tiene WhatsApp y formulario básico, pero sin chatbot ni CRM visible; falta calificar y dar seguimiento a leads.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Consultatio", "persona": "", "telefono": "", "email": "", "web": "consultatio.com.ar", "stand": "",
        "rubro": "Desarrollo Urbano", "que_vende": "Desarrollo inmobiliario (Eduardo Costantini).",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Gran desarrollador; potencial como cliente de AureaHub para CRM de ventas inmobiliarias de sus emprendimientos.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Herralum", "persona": "", "telefono": "", "email": "", "web": "herralum.com.ar", "stand": "I-2a",
        "rubro": "Herrajes y Maquinaria", "que_vende": "Herrajes para aluminio y PVC.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Fabricante directo sin intermediarios; un e-commerce B2B con login de carpinterías + CRM de seguimiento de pedidos.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Megevand Soft", "persona": "", "telefono": "", "email": "", "web": "megevand.com.ar", "stand": "I-2",
        "rubro": "Software Carpinterías", "que_vende": "OP 2.1 – Software paramétrico para carpinterías.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "SaaS de nicho; posible partnership de integración más que cliente directo de AureaHub.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Amazonas Vertical Gardens", "persona": "", "telefono": "", "email": "", "web": "amazonasverticalgardens.com", "stand": "A-6",
        "rubro": "Paisajismo", "que_vende": "Jardines verticales, techos verdes.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Nicho verde premium; un portfolio interactivo con cotizador por m² + seguimiento de proyectos vía CRM sería su upgrade.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Tope Urbano", "persona": "", "telefono": "", "email": "", "web": "topeurbano.com", "stand": "K-10",
        "rubro": "Señalización", "que_vende": "Kits autoinstalables demarcación.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Producto simple de e-commerce: un carrito optimizado + WhatsApp de seguimiento post-compra maximizaría conversión.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Lifecycle", "persona": "", "telefono": "", "email": "", "web": "", "stand": "",
        "rubro": "Software Construcción", "que_vende": "Presupuesto y costos para obra.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "SaaS de presupuestos; posible partnership/integración.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Barrisol", "persona": "", "telefono": "", "email": "", "web": "barrisol.com", "stand": "",
        "rubro": "Cielorrasos / Acústica", "que_vende": "Sistemas acústicos tensados (dist. por Wagg).",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Marca francesa distribuida por Wagg; el approach debería ser via Wagg, no directo.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Tromen", "persona": "", "telefono": "", "email": "", "web": "tromen.com", "stand": "",
        "rubro": "Calefacción", "que_vende": "Calefactores a pellet.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Marca líder calefacción; un configurador de estufas por ambiente + CRM de distribuidores escalaría su red comercial.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "Total (herramientas)", "persona": "", "telefono": "", "email": "", "web": "", "stand": "",
        "rubro": "Herramientas", "que_vende": "Herramientas línea 42V MAX.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Marca global de herramientas; posible interés en CRM para su canal de distribución argentino.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "PARSECS", "persona": "", "telefono": "", "email": "", "web": "", "stand": "",
        "rubro": "Selladores", "que_vende": "Selladores pintables.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Fabricante de selladores: un catálogo técnico digital con fichas descargables + CRM captaría leads de ferreterías.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
    {
        "tier": "TIER 3", "empresa": "PH", "persona": "", "telefono": "", "email": "", "web": "", "stand": "H-4",
        "rubro": "Construcción", "que_vende": "Productos para construcción.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Sin datos suficientes para auditar; investigar rubro y web antes de contactar.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web (Instagram)",
    },
    {
        "tier": "TIER 3", "empresa": "EVEL", "persona": "", "telefono": "", "email": "", "web": "", "stand": "",
        "rubro": "Medición", "que_vende": "Instrumentos de medición para construcción.",
        "sponsor": "", "prioridad": "", "estado": "Solo scrape",
        "dolor": "", "sistema": "", "decide": "", "objecion": "", "gancho": "",
        "auditoria": "Nicho medición; un e-commerce B2B con comparador de instrumentos + soporte técnico vía chatbot sería su diferencial.",
        "proximo": "", "contactado": "", "obs_wa": "", "fuente": "Scrape web",
    },
]

keys = ["tier", "empresa", "persona", "telefono", "email", "web", "stand", "rubro",
        "que_vende", "sponsor", "prioridad", "estado", "dolor", "sistema", "decide",
        "objecion", "gancho", "auditoria", "proximo", "contactado", "obs_wa", "fuente"]

for idx, row_dict in enumerate(rows_data):
    row = idx + 2
    for col, key in enumerate(keys, 1):
        cell = ws.cell(row=row, column=col, value=row_dict.get(key, ""))
        cell.alignment = wrap
        cell.border = thin_border

    tier = row_dict["tier"]
    if tier in tier_fills:
        ws.cell(row=row, column=1).fill = tier_fills[tier]
        ws.cell(row=row, column=1).font = tier_fonts[tier]
        ws.cell(row=row, column=1).alignment = wrap_center

col_widths = [10, 28, 30, 28, 30, 28, 8, 22, 50, 16, 16, 18, 55, 28, 24, 40, 50, 55, 35, 12, 42, 20]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:V{len(rows_data) + 1}"

# =====================================================================
# HOJA 2: RESUMEN DEBRIEF (copiada de Contactos e Insights)
# =====================================================================
ws2 = wb.create_sheet("Síntesis Debrief")

debrief = [
    ("BATEV 2026 · SÍNTESIS", ""),
    ("", ""),
    ("1 · Sensación general del rubro", ""),
    ("Mezcla de fabricantes y distribuidores de construcción y calefacción.", "Parcial"),
    ("Dato de clima económico limitado. A completar con debrief del equipo.", "Pendiente"),
    ("", ""),
    ("3 · Insights transversales", ""),
    ("Dolor más fuerte: pérdida de contexto e historial en WhatsApp, sin calificación por estadío (Parrillas/Pavir).", "Fuerte"),
    ("ERP sin capa comercial (Macoser): sumar ventas sin tocar el ERP.", "Sólido"),
    ("Empresa nueva armando base para el futuro (American Tint): demanda baja, posible caso Bardo.", "A filtrar"),
    ("Tentación de desarrollo a medida con otra agencia (Emotorart): competencia directa, URGENTE.", "Urgente"),
    ("Complejidad multimarca/multipúblico (Parrillas): un foco por marca. Ideal para Smart Tags.", "Fuerte"),
    ("", ""),
    ("5 · Competencia y marketing", ""),
    ("Competencia directa: agencias de desarrollo a medida (Emotorart).", "Ojo"),
    ("Varios tienen marketing interno o agencia. No reemplazar marketing, ordenar la máquina de ventas.", "Nota"),
    ("Permisos de redes complejos por licenciatarios (Parrillas: Bosca, Napoleón).", "Nota"),
    ("Puntos de entrada vía marketing (American Tint, Rheem). Ojo Authority: marketing no siempre decide.", "Nota"),
    ("", ""),
    ("6 · Cómo nos fue", ""),
    ("Funcionó la analogía auto hecho vs mandárselo a hacer (Emotorart). Reutilizable.", "Aprendizaje"),
    ("Resto del feedback interno: a completar con el equipo.", "Pendiente"),
    ("", ""),
    ("7 · Próximos pasos (1ra semana julio)", ""),
    ("Reuniones: Parrillas/Pavir, Macoser, GU Herrajes, Cristales Sáenz, Lindall, Asothella, Rheem.", "Agendar"),
    ("American Tint: en 2 semanas.", "Agendar"),
    ("Emotorart: URGENTE, contactar ya antes de que cierre con otra agencia.", "Urgente"),
    ("Confirmar: tarjeta parrillas, contacto Emotorart, naming Lindall, mail Rodolfo Ortiz.", "Confirmar"),
]

for idx, (detalle, peso) in enumerate(debrief):
    r = idx + 1
    ws2.cell(row=r, column=1, value=detalle).alignment = wrap
    ws2.cell(row=r, column=2, value=peso).alignment = wrap_center
    if peso == "Urgente":
        ws2.cell(row=r, column=1).font = Font(bold=True, color="CC0000")
        ws2.cell(row=r, column=2).font = Font(bold=True, color="CC0000")
    elif detalle.startswith(("1 ·", "3 ·", "5 ·", "6 ·", "7 ·", "BATEV")):
        ws2.cell(row=r, column=1).font = Font(bold=True, size=12)

ws2.column_dimensions["A"].width = 90
ws2.column_dimensions["B"].width = 15

# =====================================================================
# HOJA 3: STATS
# =====================================================================
ws3 = wb.create_sheet("Números")

tier1 = sum(1 for r in rows_data if r["tier"] == "TIER 1")
tier2 = sum(1 for r in rows_data if r["tier"] == "TIER 2")
tier3 = sum(1 for r in rows_data if r["tier"] == "TIER 3")

stats = [
    ("BATEV 2026 · Directorio Integrado AureaHub", ""),
    ("", ""),
    ("Expositores totales en BATEV", "215"),
    ("Relevados en este directorio", str(len(rows_data))),
    ("", ""),
    ("TIER 1 — Ida y vuelta real (insights, audios, tarjetas)", str(tier1)),
    ("TIER 2 — Contacto en feria (teléfono/email, WhatsApp)", str(tier2)),
    ("TIER 3 — Solo scrape web (sin contacto directo)", str(tier3)),
    ("", ""),
    ("Con auditoría web completada", str(sum(1 for r in rows_data if r["auditoria"]))),
    ("Con dolor/insight identificado", str(sum(1 for r in rows_data if r["dolor"]))),
    ("Con próximo paso definido", str(sum(1 for r in rows_data if r["proximo"]))),
    ("", ""),
    ("Datos generales BATEV 2026", ""),
    ("Fechas", "24-27 junio 2026"),
    ("Lugar", "La Rural, CABA"),
    ("Superficie", "15.000 m²"),
    ("Visitantes", "24.347"),
    ("Expositores internacionales", "58 (China, Brasil, EEUU, Turquía, Uruguay)"),
    ("Próxima edición", "BATEV 2027: 23-26 junio 2027"),
]

for idx, (dato, valor) in enumerate(stats):
    r = idx + 1
    ws3.cell(row=r, column=1, value=dato).font = Font(bold=True, size=11) if not valor else Font(size=10)
    ws3.cell(row=r, column=2, value=valor).font = Font(size=10)
    ws3.cell(row=r, column=1).alignment = wrap
    if dato.startswith("TIER"):
        ws3.cell(row=r, column=1).font = Font(bold=True, size=11)

ws3.column_dimensions["A"].width = 55
ws3.column_dimensions["B"].width = 50

output = "BATEV_2026_Directorio_Integrado_AureaHub.xlsx"
wb.save(output)
print(f"Generado: {output}")
print(f"Total empresas: {len(rows_data)}")
print(f"  TIER 1 (insights reales): {tier1}")
print(f"  TIER 2 (contacto feria): {tier2}")
print(f"  TIER 3 (solo scrape): {tier3}")
