import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Desarrollistas BA - Prospectos"

headers = [
    "Empresa",
    "Teléfono/WhatsApp",
    "Email",
    "Sitio Web",
    "Contacto (Director/CEO)",
    "Inmobiliaria Propia",
    "Usa Prometheo/Chatbot",
    "Qué desarrollan",
    "Notas"
]

data = [
    [
        "ADN Developers",
        "+54 11 4775-5158 / WA: 11-3893-3260",
        "info@adndevelopers.com.ar",
        "adndevelopers.com.ar",
        "Elbio Leandro Stoler (Fundador y Director)",
        "Sí (Nexo Propiedades)",
        "No",
        "Residencial (1-5 amb), oficinas, turístico. 20+ edificios, ~100.000 m²",
        "Palermo Hollywood. También marca DOME para torres premium"
    ],
    [
        "Land Developers Group",
        "+54 11 5263-3596 / WA: +54 9 11 2240-4473",
        "contacto@ldgcompany.com",
        "landdevelopersgroup.com",
        "Gabriel Rojchman (CEO)",
        "Sí",
        "No",
        "Townhouses exclusivos, oficinas corporativas, locales comerciales",
        "30+ años. También en Florida, USA. Dir. Comercial: Hernán Comisarenco"
    ],
    [
        "Lepore Propiedades",
        "+549 11 2249-3859 (Caballito HQ)",
        "info@lepore.com.ar",
        "lepore.com.ar",
        "Norberto Lepore (Presidente/Fundador)",
        "Sí (5 sucursales, 60+ empleados)",
        "No",
        "Residencial de categoría en CABA. También venta/alquiler usados",
        "Desde 1980. Empresa familiar 2da generación. Sucursales: Caballito, Palermo, V.Crespo"
    ],
    [
        "Vitrium Capital",
        "+54 9 11 2152 0660",
        "ventas@vitriumcapital.com",
        "vitriumcapital.com",
        "Federico Gagliardo (Fundador y CEO)",
        "Sí",
        "No",
        "Residencial, comercial, mixed-use premium",
        "Forbes Top 20 desarrollistas. Presencia en ARG, URU, PAR, USA. Proyecto Pueblo Caamaño ~US$80M"
    ],
    [
        "INARCH",
        "4774-7154 / WA: +54 9 11 5740-4253",
        "comercial@inarch.com.ar",
        "inarch.com.ar",
        "Anabela Santos (CEO), Paulo González Toledo (Co-fundador)",
        "Sí (comercialización integral)",
        "No",
        "Residencial. Edificios 5.000-7.000 m². Núñez, Palermo, Chacarita",
        "Desde 2007. Parte de Grupo GTS. ~150 empleados internos. Dir. Comercial: Jeremías González Toledo"
    ],
    [
        "Spazios",
        "0800-555-8289 / WA: +54 9 11 5217 5990",
        "info@spazios.com.ar",
        "spazios.com.ar",
        "Juan Manuel Tapiola (CEO y Fundador)",
        "Sí (Spazios Gestión Inmobiliaria)",
        "No",
        "Residencial exclusivo. 39 proyectos, 1.000+ deptos. CABA y GBA",
        "20+ años. Financiación propia 'Más Dueños' hasta 360 cuotas. Sucursales: Belgrano, Devoto, Caseros"
    ],
    [
        "Grupo ECIPSA",
        "+54 9 11 2789 4782 (BA) / +54 351 425 5525 (Cba)",
        "info@ecipsa.com",
        "ecipsa.com",
        "Jaime Garbarsky (Presidente), Walter Fuks (CEO)",
        "Sí (Natania, Natania Directa)",
        "No",
        "Residencial masivo (marca Natania), MilAires (CABA, 6.800+ unidades)",
        "Desde 1979. 9 provincias. Gte. Comercial: Ezequiel Bonomo. Proyecto MilAires US$140M+"
    ],
    [
        "Grupo Portland",
        "(+54) 11 2488-1156",
        "comercial@grupoportland.com",
        "grupoportland.com",
        "Gustavo Menayed (CEO y Fundador)",
        "No confirmado",
        "No",
        "Residencial torres, oficinas, mixed-use. L'Avenue Libertador (Zaha Hadid)",
        "20+ años. Vicente López. Expandiéndose a Miami"
    ],
    [
        "Dypsa Group",
        "+54 9 11 4468-7285",
        "info@dypsa.com",
        "dypsa.com",
        "Issel Kiperszmid (CEO y Fundador)",
        "Sí (desarrolla, construye y comercializa)",
        "No",
        "Muy diversificado: torres residenciales, shoppings, hoteles, barrios privados",
        "40+ años. Fundador de CEDU. Av. Libertador 174 P11, Vicente López. Hijo: Nicolás Kiperszmid"
    ],
    [
        "GNV Group",
        "WA: +54 9 11 3032 8205 / +54 9 11 6171 4976",
        "info@gnvgroup.com",
        "gnvgroup.com",
        "Alejandro Ginevra (Presidente y CEO)",
        "Sí (Ginevra International Realty)",
        "No",
        "Luxury mixed-use: torres residenciales, hoteles, shoppings",
        "50+ años. Madero Harbour (Puerto Madero). SLS Hotel Punta del Este"
    ],
    [
        "Newland Desarrollos",
        "11 6395-3292",
        "info@newland.com.ar",
        "newland.com.ar",
        "Ernesto Davidsohn Benet (Fundador y CEO)",
        "Sí",
        "No",
        "Residencial urbano. Canvas Barracas, VitrAUx Palermo, Smart Point",
        "20+ años. 30+ edificios. Juncal 2082, Recoleta. Presencia en ARG, URU, España"
    ],
    [
        "Vizora Desarrollos",
        "+54 11 2182-2089",
        "ventas@vizora.com.ar",
        "vizora.com.ar",
        "Milagros Brito (Fundadora). Gte. Ventas: Fabián Escanes",
        "Sí (parcial, algunos proyectos propios)",
        "No",
        "Residencial premium, comercial, cultural. Remeros Beach (Tigre), Torre Banco Macro",
        "20+ años. ZenCity (Puerto Madero) con Fernández Prieto"
    ],
    [
        "Fernández Prieto y Asoc.",
        "+54 11 4589-7792",
        "No publicado",
        "fernandezprieto.com.ar",
        "Rodrigo Fernández Prieto (CEO)",
        "Sí (también comercializa como inmobiliaria)",
        "No",
        "Torres residenciales premium (Puerto Madero), comercial, industrial, obras públicas",
        "Link Towers, Link Diamond, ZenCity, Terminal Cruceros BA. Usa Urbania 3D para ventas"
    ],
    [
        "Grupo Monarca",
        "011 2080-1800",
        "contacto@grupomonarca.com",
        "grupomonarca.com",
        "Gonzalo Monarca (CEO y Fundador)",
        "Sí (coordinación de comercialización)",
        "No",
        "Residencial. 650.000+ m² desarrollados, 20+ proyectos",
        "Desde 1997. ~500 empleados. 1ra desarrollista ARG con ISO 9001. Arias 1639, CABA"
    ],
    [
        "TGLT (ahora GCDI)",
        "WA: +54 9 11 2453 5448 / Tel: +54 11 5252 5050",
        "ventas@tglt.com",
        "tglt.com",
        "Federico Weil (Fundador, ya no activo como CEO)",
        "Sí (ciclo completo)",
        "No",
        "Residencial y mixed-use. Marcas: Forum, Astor, Metra, Venice. 570.000+ m²",
        "Cotiza en BYMA (TGLT.BA). Miñones 2177, CABA"
    ],
    [
        "Consultatio",
        "+54 11 4318-8000",
        "info@consultatio.com.ar",
        "consultatio.com.ar",
        "Eduardo Costantini (Fundador/Presidente), Gonzalo de la Serna (CEO)",
        "Sí (Consultatio Real Estate)",
        "No",
        "Urbanizaciones (Nordelta), residencial, torres corporativas",
        "Desde 1991. 501-1.000 empleados. Adquirió 51% de Argencons en 2025. ARG, URU, USA"
    ],
    [
        "IRSA Propiedades",
        "0800-222-2299",
        "Formulario web",
        "irsa.com.ar",
        "Eduardo Elsztain (Presidente)",
        "Sí (comercialización interna)",
        "No",
        "Shoppings, oficinas, hoteles, residencial (Ramblas del Plata)",
        "Mayor empresa inmobiliaria cotizante de ARG. BYMA + NYSE. Desde 1943"
    ],
    [
        "Grupo Farallón",
        "WA: +54 9 11 6815-8480",
        "No publicado",
        "grupofarallon.com",
        "Gerónimo Gutiérrez de Barrio (Presidente)",
        "No confirmado (vende proyectos propios)",
        "No",
        "Barrios privados, countries, centros comerciales, vivienda social",
        "Desde 1985. Ruta Panamericana Km 49.5, Pilar. Pioneros del barrio privado en GBA"
    ],
    [
        "Fortune International (Defortuna)",
        "BA: +5411 7700-1010 / Miami: (305) 400-6393",
        "info@fir.com",
        "fortuneintlgroup.com",
        "Edgardo Defortuna (Presidente, CEO y Fundador)",
        "Sí (Christie's International RE)",
        "No",
        "Luxury residencial. Principalmente South Florida/Miami. BA como oficina de ventas",
        "Argentino radicado en Miami. 18 oficinas mundiales. ~1.000 asociados"
    ],
    [
        "Grupo Briones",
        "+54 9 11 2468-2070",
        "No publicado",
        "grupobriones.com.ar",
        "No identificado",
        "Sí (comercialización propia)",
        "No",
        "Residencial en CABA",
        "Encontrado en investigación previa. Sin más datos públicos"
    ],
    [
        "Alto Grande / M&M Propiedades",
        "+11 5272-0800",
        "No publicado",
        "No encontrado",
        "No identificado",
        "Sí (M&M Propiedades es su inmobiliaria)",
        "No",
        "Residencial en CABA",
        "Tiene inmobiliaria propia M&M Propiedades"
    ],
    [
        "Big Ben",
        "+54 (11) 4483-4566",
        "No publicado",
        "No encontrado",
        "Andrés Castellano (CEO)",
        "No confirmado",
        "No",
        "Residencial en CABA",
        ""
    ],
    [
        "Grow Desarrollos",
        "No encontrado",
        "No encontrado",
        "No encontrado",
        "No identificado",
        "Sí (tiene inmobiliaria propia)",
        "No",
        "Residencial en Buenos Aires",
        "Referencia del usuario. Datos públicos limitados"
    ],
]

header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

for row_num, row_data in enumerate(data, 2):
    for col_num, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_num, column=col_num, value=value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = thin_border
        cell.font = Font(name="Calibri", size=10)

col_widths = [28, 35, 30, 25, 40, 32, 18, 50, 55]
for i, width in enumerate(col_widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

ws.auto_filter.ref = f"A1:I{len(data) + 1}"
ws.freeze_panes = "A2"

excluded_ws = wb.create_sheet("Excluidos (Clientes Prometheo)")
excluded_headers = ["Empresa", "Motivo de Exclusión", "Sitio Web"]
excluded_data = [
    ["Argencons / Quartier", "Cliente actual de Prometheo (confirmado en prometheo.ai)", "argencons.com"],
    ["NorthBaires", "Cliente actual de Prometheo (confirmado en prometheo.ai)", "northbaires.com"],
    ["Crea Urbana", "Cliente actual de Prometheo (confirmado en prometheo.ai)", "creaurbana.com.ar"],
    ["G&D Developers", "Cliente actual de Prometheo (confirmado por usuario)", "gddevelopers.com"],
    ["TKVA", "Cliente actual de Prometheo (confirmado por usuario)", ""],
    ["EDFAN Constructora/Real Estate", "Cliente actual de Prometheo (confirmado por usuario)", ""],
]

excluded_fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid")
for col_num, header in enumerate(excluded_headers, 1):
    cell = excluded_ws.cell(row=1, column=col_num, value=header)
    cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    cell.fill = excluded_fill
    cell.alignment = header_alignment
    cell.border = thin_border

for row_num, row_data in enumerate(excluded_data, 2):
    for col_num, value in enumerate(row_data, 1):
        cell = excluded_ws.cell(row=row_num, column=col_num, value=value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = thin_border
        cell.font = Font(name="Calibri", size=10)

excluded_ws.column_dimensions["A"].width = 30
excluded_ws.column_dimensions["B"].width = 55
excluded_ws.column_dimensions["C"].width = 25

output_path = "/home/user/AUREA-Hub/Desarrollistas_BA_Prospectos.xlsx"
wb.save(output_path)
print(f"Excel generado: {output_path}")
print(f"Total prospectos: {len(data)}")
print(f"Total excluidos: {len(excluded_data)}")
