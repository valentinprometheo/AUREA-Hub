import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Expositores BATEV 2026"

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
wrap = Alignment(wrap_text=True, vertical="top")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

sponsor_fills = {
    "Main Sponsor": PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
    "Sponsor Diamond": PatternFill(start_color="B9F2FF", end_color="B9F2FF", fill_type="solid"),
    "Sponsor Platinum": PatternFill(start_color="E5E4E2", end_color="E5E4E2", fill_type="solid"),
    "Sponsor Gold": PatternFill(start_color="FFEAAD", end_color="FFEAAD", fill_type="solid"),
}

headers = [
    "Empresa", "Stand", "Rubro / Categoría", "Productos / Servicios",
    "Nivel Sponsor", "Web", "Fuente",
]

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = thin_border

expositores = [
    # Main Sponsors
    ("Clarín ARQ", "", "Medios / Arquitectura", "Suplemento de arquitectura del diario Clarín", "Main Sponsor", "", "batev.com.ar/expositores"),
    ("Muchtek", "220 m²", "Aberturas PVC", "Sistemas de perfiles de PVC para aberturas: línea Evolution, corrediza Advance, línea Efficient, In&Out Design", "Main Sponsor", "muchtek.com", "batev.com.ar + informeconstruccion.com"),

    # Diamond Sponsors
    ("Motorarg", "", "Herramientas / Maquinaria", "Herramientas y equipos para construcción", "Sponsor Diamond", "motorarg.com", "batev.com.ar/expositores"),
    ("Pavir S.A.", "E-24", "Aberturas", "Aberturas de alta prestación en acero, PVC y aluminio; puertas pivotantes con marco de aluminio", "Sponsor Diamond", "pavir.com.ar", "batev.com.ar + artículo novedades"),
    ("Sinteplast", "", "Pinturas y Revestimientos", "Pinturas, impermeabilizantes, revestimientos para construcción", "Sponsor Diamond", "sinteplast.com.ar", "norteenlinea.com"),
    ("Vann", "", "Griferías / Sanitarios", "Griferías y accesorios para baño y cocina", "Sponsor Diamond", "vann.com.ar", "batev.com.ar/expositores"),
    ("Fenstech", "", "Aberturas PVC", "Distribuidor oficial de perfiles PVC VERATEC del Grupo ASAS (Europa)", "Sponsor Diamond", "fenstech.com", "informeconstruccion.com"),

    # Platinum Sponsors
    ("Aluwind", "", "Aberturas Aluminio", "Sistemas de aberturas en aluminio", "Sponsor Platinum", "aluwind.com.ar", "batev.com.ar/expositores"),
    ("American Tint", "", "Films / Láminas", "Láminas de seguridad, control solar y decorativas", "Sponsor Platinum", "americantint.com.ar", "batev.com.ar/expositores"),
    ("A.D. Barbieri", "", "Construcción en seco", "Sistemas constructivos, ingeniería, Consulsteel, Barbieri|Deceuninck", "Sponsor Platinum", "barbieri.com.ar", "batev.com.ar + informeconstruccion.com"),
    ("Patagonia Flooring", "", "Pisos", "Pisos de madera, vinílicos y laminados", "Sponsor Platinum", "patagoniafloorings.com", "batev.com.ar/expositores"),
    ("Wagg", "E-21", "Arquitectura Textil", "Cielorrasos tensados Barrisol, cubiertas y fachadas textiles, membranas flexibles", "Sponsor Platinum", "wagg.com.ar", "batev.com.ar + artículo novedades"),

    # Gold Sponsors
    ("Aluplexa", "", "Aberturas / Persianas", "Persianas de grandes dimensiones, mosquiteras plisadas", "Sponsor Gold", "aluplexa.com", "batev.com.ar + informeconstruccion.com"),
    ("Ensa Ascensores", "", "Ascensores", "Ascensores residenciales PVE52, PVE37, PVE30 (instalación en 1 día)", "Sponsor Gold", "ensaascensores.com", "informeconstruccion.com"),
    ("Grupo LTN", "", "Soluciones constructivas", "Soluciones constructivas integrales", "Sponsor Gold", "", "mercado.com.ar"),
    ("Puertas Brandsen", "", "Puertas", "Puertas para construcción residencial y comercial", "Sponsor Gold", "", "batev.com.ar/expositores"),
    ("Home Access Argentina", "", "Accesibilidad", "Soluciones de acceso para el hogar", "Sponsor Gold", "", "batev.com.ar/expositores"),

    # Expositores (de artículos de novedades del sitio oficial)
    ("Liquitech", "A-26", "Pinturas / Impermeabilizantes", "AquaEpoxi (hidroesmalte epoxi renovación azulejos); Cauchogoma (impermeabilizantes: Rojo Teja, Verde Césped)", "", "liquitech.com.ar", "batev.com.ar/novedades"),
    ("Coef Track", "A-17", "Software Construcción", "Sistema de gestión de personal de obra: capacitación, motivación y control en tiempo real", "", "coeftrack.com", "batev.com.ar/novedades"),
    ("Finnegans GO", "C-28", "Software Construcción", "Herramienta en la nube para gestión de obras: presupuestos, materiales, contratos", "", "finneg.com/ar/", "batev.com.ar/novedades"),
    ("Herralum", "I-2a", "Herrajes y Maquinaria", "Herrajes y maquinarias para aluminio y PVC: fallebas, cierres, sistemas multi-punto, fresadoras", "", "herralum.com.ar", "batev.com.ar/novedades"),
    ("Megevand Soft", "I-2", "Software Carpinterías", "OP 2.1 – Software de cálculo paramétrico para carpinterías de aluminio y PVC", "", "megevand.com.ar", "batev.com.ar/novedades"),
    ("Amazonas Vertical Gardens", "A-6", "Paisajismo", "Jardines verticales hidropónicos, techos verdes, mantenimiento corporativo", "", "amazonasverticalgardens.com", "batev.com.ar/novedades"),
    ("Tope Urbano", "K-10", "Señalización", "Kits autoinstalables para demarcación de espacios sin pintura", "", "topeurbano.com", "batev.com.ar/novedades"),
    ("ANSAL", "C-35", "Climatización", "Sistemas GREE: aire acondicionado + piso radiante en un equipo; calderas a gas, VRF, splits", "", "ansal.com.ar", "batev.com.ar/novedades"),
    ("IDERO", "F-7a", "Construcción Modular", "Lodges premium modulares (18-67 m²), fabricación en fábrica, montaje en 1 día", "", "ideroarquitectura.com.ar", "batev.com.ar/novedades"),
    ("Almapiedra", "H-11", "Restauración / Nanotecnología", "Restauración patrimonial, distribuidores exclusivos Tecnan (nanotecnología protección piedra/hormigón)", "", "almapiedra.com.ar", "batev.com.ar/novedades"),
    ("Cambre", "E-11a", "Electricidad / Domótica", "Soluciones de diseño, conectividad, Cambre Live Home (automatización inteligente)", "", "cambre.com.ar", "batev.com.ar/novedades"),
    ("Rheem", "D-16", "Agua Caliente", "Calefones, calderas SmartHeatDuo, termotanques residenciales/comerciales, híbridos", "", "rheem.com.ar", "batev.com.ar/novedades"),

    # Expositores de otras fuentes (Instagram, prensa)
    ("Lifecycle", "", "Software Construcción", "Herramientas de presupuesto y costos para obra", "", "", "mercado.com.ar"),
    ("Elektrim", "", "Tratamiento de Agua", "Ablandadores de agua", "", "", "mercado.com.ar"),
    ("Barrisol", "", "Cielorrasos / Acústica", "Sistemas acústicos y revestimientos tensados (distribuido por Wagg)", "", "barrisol.com", "mercado.com.ar"),
    ("Tromen", "", "Calefacción", "Calefactores a pellet", "", "tromen.com", "mercado.com.ar"),
    ("Total (herramientas)", "", "Herramientas", "Herramientas línea 42V MAX", "", "", "pyfsitio.com"),
    ("PARSECS", "", "Selladores", "Selladores pintables para construcción", "", "", "pyfsitio.com"),
    ("Brimax", "H-1", "Construcción", "Productos para la industria de la construcción", "", "", "Instagram @brimax"),
    ("PH", "H-4", "Construcción", "Productos para construcción", "", "", "Instagram"),
    ("Consultatio", "", "Desarrollo Urbano", "Desarrollo inmobiliario y urbano (Presidente: Eduardo Costantini)", "", "consultatio.com.ar", "todoenunclick.com"),

    # Expositores de artículos novedades (deducidos de URLs)
    ("EVEL", "", "Medición", "Instrumentos de medición argentina para la construcción e instalación", "", "", "batev.com.ar/novedades (URL)"),
]

for idx, (empresa, stand, rubro, productos, sponsor, web, fuente) in enumerate(expositores):
    row = idx + 2
    vals = [empresa, stand, rubro, productos, sponsor, web, fuente]
    for col, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.alignment = wrap
        cell.border = thin_border

    if sponsor in sponsor_fills:
        for col in range(1, 8):
            ws.cell(row=row, column=col).fill = sponsor_fills[sponsor]

col_widths = [28, 10, 28, 60, 18, 30, 32]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:G{len(expositores) + 1}"

# --- HOJA 2: Resumen por categoría ---
ws2 = wb.create_sheet("Resumen por Categoría")
ws2.cell(row=1, column=1, value="Expositores BATEV 2026 por Categoría").font = Font(bold=True, size=14)

from collections import Counter
cats = Counter(e[2] for e in expositores)

ws2.cell(row=3, column=1, value="Categoría").font = Font(bold=True)
ws2.cell(row=3, column=2, value="Cantidad").font = Font(bold=True)
ws2.cell(row=3, column=3, value="Empresas").font = Font(bold=True)

for idx, (cat, count) in enumerate(cats.most_common()):
    r = 4 + idx
    ws2.cell(row=r, column=1, value=cat)
    ws2.cell(row=r, column=2, value=count)
    empresas_cat = [e[0] for e in expositores if e[2] == cat]
    ws2.cell(row=r, column=3, value=", ".join(empresas_cat))
    for col in range(1, 4):
        ws2.cell(row=r, column=col).alignment = wrap
        ws2.cell(row=r, column=col).border = thin_border

ws2.column_dimensions["A"].width = 30
ws2.column_dimensions["B"].width = 12
ws2.column_dimensions["C"].width = 60

# --- HOJA 3: Stats ---
ws3 = wb.create_sheet("Datos Generales BATEV")

stats = [
    ("Dato", "Valor"),
    ("Nombre", "BATEV – Exposición Internacional de la Construcción y la Vivienda"),
    ("Edición", "31ª edición"),
    ("Fechas", "24-27 junio 2026"),
    ("Horario", "14 a 20 hs"),
    ("Lugar", "La Rural, Av. Sarmiento 2704, Palermo, CABA"),
    ("Superficie", "15.000 m²"),
    ("Expositores totales", "215 empresas (58 internacionales: China, Brasil, EEUU, Turquía, Uruguay)"),
    ("Visitantes", "24.347"),
    ("Workshops", "25"),
    ("Organizadores", "CEDU+AEV, CAMARCO, EFCA S.A."),
    ("EFCA contacto", "Av. de Mayo 605, 11° Piso A, CABA, C1084AAB"),
    ("Presidente EFCA", "Gabriel Pascual"),
    ("Presidente CEDU+AEV", "Damián Tabakman"),
    ("Vicepresidente CEDU+AEV", "Carlos Spina"),
    ("Expositores relevados en este archivo", f"{len(expositores)} de 215"),
    ("Próxima edición", "BATEV 2027: 23-26 junio 2027, La Rural"),
    ("Fuentes", "batev.com.ar, informeconstruccion.com, mercado.com.ar, todoenunclick.com, pyfsitio.com, Instagram"),
]

for idx, (dato, valor) in enumerate(stats):
    r = idx + 1
    ws3.cell(row=r, column=1, value=dato).font = Font(bold=True) if idx == 0 else Font(bold=True, size=10)
    ws3.cell(row=r, column=2, value=valor)
    ws3.cell(row=r, column=1).alignment = wrap
    ws3.cell(row=r, column=2).alignment = wrap

ws3.column_dimensions["A"].width = 30
ws3.column_dimensions["B"].width = 70

output = "Expositores_BATEV_2026.xlsx"
wb.save(output)
print(f"Generado: {output}")
print(f"Total expositores relevados: {len(expositores)}")
